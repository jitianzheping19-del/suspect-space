"""Word files contain literal code, with exact line breaks and no HTML import."""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

directory = Path(sys.argv[1])
for filename in ("code.js", "index.html"):
    source = (directory / filename).read_text(encoding="utf-8")
    document = Document()
    document.core_properties.title = filename
    document.core_properties.subject = "SPACE SUSPECT / Google Apps Script"
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.55)
    style = document.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(8)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1
    # One paragraph preserves the source exactly, including its final newline.
    paragraph = document.add_paragraph()
    paragraph.add_run(source)
    props = paragraph._p.get_or_add_pPr()
    no_wrap = OxmlElement("w:wordWrap")
    no_wrap.set(qn("w:val"), "0")
    props.append(no_wrap)
    destination = directory / (filename + ".docx")
    document.save(destination)
    extracted = Document(destination).paragraphs[0].text
    if extracted != source:
        raise ValueError("Word source differs from original: " + filename)
    print(f"Created {destination.name}: {destination.stat().st_size:,} bytes, exact code round trip")
